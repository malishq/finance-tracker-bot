from __future__ import annotations

import csv
import io
from datetime import date
from typing import Iterable, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import INCOME, Entry

HEADERS = ["Дата", "Тип", "Категория", "Описание", "Сумма"]


def _kind_label(kind: str) -> str:
    return "Доход" if kind == INCOME else "Расход"


def _rows(entries: Iterable[Entry]) -> List[list]:
    rows = []
    for entry in sorted(entries, key=lambda e: (e.entry_date, e.id)):
        signed = entry.amount if entry.kind == INCOME else -entry.amount
        rows.append([entry.entry_date, _kind_label(entry.kind), entry.category, entry.description, signed / 100])
    return rows


def entries_to_csv(entries: Iterable[Entry]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow(HEADERS)
    for row in _rows(entries):
        amount = f"{row[4]:.2f}".replace(".", ",")
        writer.writerow([row[0].strftime("%d.%m.%Y"), row[1], row[2], row[3], amount])
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def entries_to_xlsx(entries: Iterable[Entry], currency: str = "₽") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Операции"
    sheet.append(HEADERS)
    header_fill = PatternFill("solid", fgColor="1F2937")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    rows = _rows(entries)
    for row in rows:
        sheet.append(row)
    number_format = f'#,##0.00 "{currency}"'
    for index in range(2, len(rows) + 2):
        sheet.cell(row=index, column=1).number_format = "DD.MM.YYYY"
        sheet.cell(row=index, column=5).number_format = number_format
    total_row = len(rows) + 2
    income = sum(r[4] for r in rows if r[4] > 0)
    expense = -sum(r[4] for r in rows if r[4] < 0)
    sheet.cell(row=total_row + 1, column=4, value="Доходы").font = Font(bold=True)
    sheet.cell(row=total_row + 1, column=5, value=income).number_format = number_format
    sheet.cell(row=total_row + 2, column=4, value="Расходы").font = Font(bold=True)
    sheet.cell(row=total_row + 2, column=5, value=expense).number_format = number_format
    sheet.cell(row=total_row + 3, column=4, value="Баланс").font = Font(bold=True)
    balance_cell = sheet.cell(row=total_row + 3, column=5, value=income - expense)
    balance_cell.number_format = number_format
    balance_cell.font = Font(bold=True)
    widths = [12, 10, 22, 36, 16]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_filename(start: date, end: date, extension: str) -> str:
    if start.year == end.year and start.month == end.month:
        return f"finance_{start:%Y-%m}.{extension}"
    return f"finance_{start:%Y-%m-%d}_{end:%Y-%m-%d}.{extension}"
